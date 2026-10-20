#!/usr/bin/env python3
"""Extrae la LISTA DE PRECIOS (xlsx) y genera data/catalog_seed.json.

Clasifica cada artículo en un tipo de componente del dominio (CPU, MOBO, RAM,
PSU, GABINETE, ALMACENAMIENTO, ENFRIAMIENTO AIRE/LÍQUIDO, GPU, MONITOR,
TECLADO, MOUSE, KIT, TARJETA EXPANSIÓN, OTRO) y extrae especificaciones
relevantes (watts de PSU, PSU incluida en gabinete, disipador incluido en CPU,
DDR/capacidad de RAM, tipo SSD/HDD, etc.).

Uso:
    python3 scripts/extract_catalog.py [ruta/al/xlsx] [-o data/catalog_seed.json]

El tipo de cambio (T.C) se toma de la fila 2 de la hoja. La regla de negocio es
redondearlo hacia arriba al entero superior (18.52 -> 19) para convertir USD a MXN.
"""
import argparse
import json
import re
import sys
import uuid
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl

NAMESPACE = uuid.UUID("6ba7b811-9dad-11d1-80b4-00c04fd430c8")  # uuid5 ns:URL

WAREHOUSES = ["a_cen", "ave_2", "cedis", "cmal", "gdl"]


def stable_id(sku: str, description: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f"{sku}|{description}"))


def money(amount, currency: str):
    return {
        "amount": str(Decimal(str(amount)).quantize(Decimal("0.01"))),
        "currency": "USD" if str(currency).strip().upper() == "USD" else "MXP",
    }


def classify(department: str, desc: str, sku: str):
    """Devuelve (kind, specs) para un artículo."""
    d = desc.upper()
    specs = {}

    if department == "PROCESADORES":
        m = re.search(r"\b(AM4|AM5|LGA ?\d{3,4})\b", d)
        if m:
            specs["socket"] = m.group(1).replace(" ", "")
        if "SIN DISIPADOR" in d or "NO INCLUYE DISIPADOR" in d:
            specs["includes_cooler"] = False
        elif sku.upper().endswith("WOF"):
            specs["includes_cooler"] = False
        elif "CON DISIPADOR" in d or "INCLUYE DISIPADOR" in d or sku.upper().endswith("BOX"):
            specs["includes_cooler"] = True
        return "cpu", specs

    if department == "MOTHER BOARD":
        m = re.search(r"\b(AM4|AM5|LGA ?\d{3,4})\b", d)
        if m:
            specs["socket"] = m.group(1).replace(" ", "")
        ff = "MINI ITX" if "MINI ITX" in d else ("MICRO ATX" if "MICRO ATX" in d or "MATX" in d else ("ATX" if " ATX" in d else None))
        if ff:
            specs["form_factor"] = ff
        return "motherboard", specs

    if department == "MEMORIA RAM":
        m = re.search(r"\b(DDR[345])\b", d)
        if m:
            specs["ddr"] = m.group(1)
        m = re.search(r"(\d+)\s*(?:GB|G )", d)
        if m:
            specs["capacity_gb"] = int(m.group(1))
        specs["ram_form"] = "SODIMM" if "SODIMM" in d else "DIMM"
        return "ram", specs

    if department == "TARJETA DE VIDEO":
        m = re.search(r"(\d+GB)\s*(G?DDR\d\w*)", d)
        if m:
            specs["vram"] = f"{m.group(1)} {m.group(2)}"
        return "gpu", specs

    if department == "MONITORES":
        m = re.search(r"\b(\d{2}(?:\.\d)?)\s*(?:IN|P\b|PLASMA)", d)
        if m:
            specs["inches"] = m.group(1)
        return "monitor", specs

    if department == "SSD INTERNO":
        specs["storage_kind"] = "ssd"
        specs["internal"] = True
        return "storage", specs

    if department in ("DD INTERNO", "DISCO DURO") or (department == "DISCO DURO" and "INTERNO" in d):
        specs["storage_kind"] = "hdd"
        specs["internal"] = True
        return "storage", specs

    if department == "ENFRIAMIENTO":
        if "LIQUIDO" in d or "LIQUIDA" in d or " AIO" in d:
            return "cooler_liquid", specs
        if "BASE ENFRIADORA" in d or "LAPTOP" in d:
            return "other", specs
        return "cooler_air", specs

    if department == "PERIFERICOS":
        if "MOUSE" in d and "TECLADO" in d:
            return "keyboard_mouse_kit", specs
        if "MOUSE" in d:
            return "mouse", specs
        if "TECLADO" in d:
            return "keyboard", specs
        return "other", specs

    if department == "REDES":
        return "expansion_card", {"card_type": "red"}

    # COMPONENTES mezcla fuentes de poder y gabinetes
    if department == "COMPONENTES":
        if "FUENTE DE PODER" in d:
            m = re.search(r"\b(\d{3,4})\s*W", d)
            if m:
                specs["wattage"] = int(m.group(1))
            return "psu", specs
        if "GABINETE" in d or "CASE" in d:
            if "SIN FUENTE" in d:
                specs["includes_psu"] = False
            elif "CON FUENTE" in d:
                specs["includes_psu"] = True
                m = re.search(r"FUENTE\s*(?:DE PODER\s*)?(\d{3,4})\s*W", d)
                if m:
                    specs["psu_watts"] = int(m.group(1))
            else:
                specs["includes_psu"] = False
            return "case", specs
        return "other", specs

    # ACCESORIOS: una tarjeta de red PCI es una tarjeta de expansión
    if department == "ACCESORIOS":
        if "TARJETA DE RED" in d and ("PCI" in d or "PCIE" in d):
            return "expansion_card", {"card_type": "red"}
        return "other", specs

    return "other", specs


def capacity_from_desc(desc: str):
    m = re.search(r"\b(\d+(?:\.\d+)?)\s*(TB|GB)\b", desc.upper())
    if not m:
        return None
    val = Decimal(m.group(1))
    unit = m.group(2)
    return {"raw": f"{m.group(1)}{unit}", "gb": int(val * 1000) if unit == "TB" else int(val)}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx", nargs="?", default="/Users/harunoyukki/Downloads/LISTA DE PRECIOS  15-10-26.xlsx")
    ap.add_argument("-o", "--out", default="data/catalog_seed.json")
    args = ap.parse_args()

    src = Path(args.xlsx)
    if not src.exists():
        sys.exit(f"No existe: {src}")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))

    # Fila 2: fecha y T.C
    exchange_rate = None
    list_date = None
    for cell in rows[1]:
        s = str(cell) if cell is not None else ""
        m = re.search(r"T\.?C\.?\s*([\d]+(?:\.\d+)?)", s.upper())
        if m:
            exchange_rate = m.group(1)
        if isinstance(cell, datetime):
            list_date = cell.date().isoformat()
    if exchange_rate is None:
        sys.exit("No se encontró el T.C en la fila 2")

    components = []
    seen = set()
    stats = {}
    for r in rows[3:]:
        sku = str(r[0]).strip() if r[0] is not None else ""
        desc = str(r[2]).strip() if r[2] is not None else ""
        if not sku and not desc:
            continue
        key = (sku, desc)
        if key in seen:
            continue
        seen.add(key)
        department = str(r[3]).strip() if r[3] is not None else ""
        currency = str(r[4]).strip() if r[4] is not None else "USD"
        brand = str(r[5]).strip() if r[5] is not None else "SIN MARCA"
        promo = str(r[6]).strip().upper() == "SI" if r[6] is not None else False
        cost = Decimal(str(r[12])) if r[12] is not None else Decimal("0")

        kind, specs = classify(department, desc, sku)
        if kind == "storage":
            cap = capacity_from_desc(desc)
            if cap:
                specs["capacity_raw"] = cap["raw"]
                specs["capacity_gb"] = cap["gb"]

        comp = {
            "id": stable_id(sku, desc),
            "sku": sku,
            "barcode": str(int(r[1])) if isinstance(r[1], (int, float)) and r[1] else (str(r[1]).strip() if r[1] else None),
            "description": desc,
            "brand": brand,
            "department": department,
            "kind": kind,
            "condition": "nuevo",
            "price": money(cost, currency),
            "on_promotion": promo,
            "stock": {w: int(r[7 + i] or 0) for i, w in enumerate(WAREHOUSES)},
            "source": "lista",
            "specs": specs,
        }
        components.append(comp)
        stats[kind] = stats.get(kind, 0) + 1

    seed = {
        "source_file": src.name,
        "list_date": list_date or date.today().isoformat(),
        "exchange_rate": exchange_rate,
        "exchange_rule": "USD→MXP: tipo de cambio redondeado hacia arriba al entero superior (18.52 → 19)",
        "components": components,
    }

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(seed, ensure_ascii=False, indent=1), encoding="utf-8")

    print(f"Generado {out}: {len(components)} componentes")
    efectivo = Decimal(exchange_rate).quantize(Decimal("1"), rounding="ROUND_CEILING")
    print(f"  T.C: {exchange_rate} (efectivo: {efectivo})")
    for k, v in sorted(stats.items(), key=lambda x: -x[1]):
        print(f"  {v:4d}  {k}")


if __name__ == "__main__":
    main()
